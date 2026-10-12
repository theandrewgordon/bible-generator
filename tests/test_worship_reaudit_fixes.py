"""Regression tests for the 2026-10-11 Worship re-audit fixes."""

import re
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

import app

ROOT = Path(__file__).resolve().parents[1]


def _deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def _slide_texts(prs) -> list[str]:
    return [
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]


class CreditLineTests(unittest.TestCase):
    def test_full_credit_line(self):
        song = {"type": "song", "copyright_notice": "Brian Johnson · © 2021 Bethel Music", "ccli_song_number": "7181 1150"}
        self.assertEqual(
            app._worship_credit_line(song, "1234567"),
            "Brian Johnson · © 2021 Bethel Music · CCLI Song #71811150 · CCLI License #1234567",
        )

    def test_public_domain_and_unlabeled_items(self):
        self.assertEqual(app._worship_credit_line({"type": "song", "copyright_notice": "Public Domain"}, "1"), "Public domain")
        self.assertEqual(app._worship_credit_line({"type": "song"}, "1234567"), "")
        self.assertEqual(app._worship_credit_line({"type": "scripture", "copyright_notice": "x"}, "1"), "")


class DeckTests(unittest.TestCase):
    def test_title_slide_shows_credit_not_key(self):
        prs = _deck()
        app.create_divider_slide(prs, "Holy Forever", "Chris Tomlin", "D", "song", None, credit="© 2021 X · CCLI License #1")
        texts = _slide_texts(prs)
        self.assertIn("© 2021 X · CCLI License #1", texts)
        self.assertFalse(any("Key:" in text for text in texts))

    def test_section_label_moves_to_presenter_notes(self):
        prs = _deck()
        app.create_content_slide(prs, ["Line one"], "song", None, part_label="Verse 1")
        self.assertNotIn("VERSE 1", _slide_texts(prs))
        self.assertEqual(prs.slides[0].notes_slide.notes_text_frame.text, "VERSE 1")

    def test_one_font_size_per_section(self):
        long_lines = ["Now as they were eating, Jesus took bread, and after blessing it broke it",
                      "and gave it to the disciples, and said, Take, eat; this is my body.",
                      "And he took a cup,", "and when he had given thanks",
                      "he gave it to them, saying,", "Drink of it, all of you."]
        song = {"type": "scripture", "parts": {"reading": long_lines}, "arrangement": ["reading"]}
        sizes = {chunk["font_size"] for _, chunk in app._iter_worship_lyric_chunks(app.normalize_worship_song(song))}
        self.assertEqual(len(sizes), 1)


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.presenter = (ROOT / "templates" / "worship_live_presenter.html").read_text(encoding="utf-8")
        self.worship = (ROOT / "templates" / "worship.html").read_text(encoding="utf-8")

    def test_leader_codes_are_hidden_from_the_room(self):
        self.assertIn('<div class="wl-leader-codes" id="wl-leader-codes" hidden>', self.presenter)
        self.assertIn("{showLeaderCodes(false);setup.hidden=!setup.hidden;", self.presenter)  # Q key
        self.assertNotIn("('Key: ' ~ slide.key)", self.presenter)

    def test_item_menu_has_no_dropdown_chevron(self):
        self.assertIn(".ws-item-menu>summary:after{content:none}", self.worship)

    def test_quick_scripture_closes_on_escape(self):
        self.assertIn("event.key === 'Escape' && quickScriptureModal.classList.contains('is-open')", self.worship)
        self.assertNotIn('class="ws-scripture-modes" role="tablist"', self.worship)


class PlannerPayloadTests(unittest.TestCase):
    def test_library_js_drops_server_only_fields(self):
        song = {
            "id": "a", "title": "A", "type": "song", "parts": {"v1": ["x"]},
            "validation": {"status": "verified", "sources": ["big"] * 50}, "review": {"status": "approved"},
            "presentation_slides": [{"hidden": False, "image_path": "p.png", "text": "long"}],
        }
        slim = app._worship_library_for_js([song])[0]
        self.assertEqual(set(slim), {"id", "title", "type", "parts", "presentation_slides"})
        self.assertEqual(slim["presentation_slides"], [{"hidden": False}])


class SimpleModeReviewTests(unittest.TestCase):
    def test_simple_mode_skips_validation_nags(self):
        source = (ROOT / "app.py").read_text(encoding="utf-8")
        block = source[source.index("def worship_deck_review():"):source.index("def worship_deck_review_complete")]
        self.assertRegex(block, re.compile(r"if simple_mode:\s+# Simple mode hides"))


if __name__ == "__main__":
    unittest.main()
