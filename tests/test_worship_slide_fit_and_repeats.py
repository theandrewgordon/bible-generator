from pptx import Presentation

import app


def _song(arrangement):
    return app.normalize_worship_song({
        "id": "test-song",
        "title": "Test Song",
        "parts": {
            "verse1": ["Verse line one", "Verse line two"],
            "chorus": ["Chorus line one", "Chorus line two"],
        },
        "arrangement": arrangement,
    })


def test_live_chunks_collapse_adjacent_repeat_into_cue():
    chunks = list(app._iter_worship_lyric_chunks(_song(["verse1", "chorus", "chorus"])))
    assert [chunk["repeat_count"] for _, chunk in chunks] == [1, 2]


def test_pptx_emits_one_slide_per_sung_repeat_with_notes():
    prs = Presentation()
    added = app.add_worship_lyric_slides(prs, _song(["verse1", "chorus", "chorus"]))
    assert added == 3
    notes = [slide.notes_slide.notes_text_frame.text for slide in prs.slides]
    assert notes[0] == "VERSE 1"
    assert notes[1:] == ["CHORUS (1 OF 2)", "CHORUS (2 OF 2)"]


def test_width_uses_real_bold_metrics_wide_vs_narrow():
    # Same character count, very different printed width.
    assert app._lyric_width_units("WWWWWWWWWW") > 2.5 * app._lyric_width_units("iiiiiiiiii")


def test_short_lines_keep_preferred_size_long_caps_shrink():
    assert app._lyric_font_size_for_width(["Amazing grace"], 48) == 48
    long_caps = ["HOLY HOLY HOLY LORD GOD ALMIGHTY WHO WAS"]
    size = app._lyric_font_size_for_width(long_caps, 48)
    assert 32 <= size < 48
    # The chosen size actually fits the ~874pt text area.
    assert app._lyric_width_units(long_caps[0]) * size <= 874


def test_non_latin_text_falls_back_without_error():
    assert app._lyric_width_units("耶稣爱我") > 0


def test_blank_line_in_editor_text_becomes_slide_break():
    lines, breaks = app._part_text_to_lines_and_breaks("\nA\nB\n\n\nC\nD\n\n")
    assert lines == ["A", "B", "C", "D"]
    assert breaks == [2]


def test_normalize_keeps_only_valid_breaks_and_parts_stay_clean():
    song = app.normalize_worship_song({
        "title": "T",
        "parts": {"Chorus": ["A", "B", "C", "D"]},
        "slide_breaks": {"Chorus": [2, 0, 4, 99, "x", 2], "ghost": [1]},
    })
    assert song["parts"]["chorus"] == ["A", "B", "C", "D"]  # no blank lines stored in lyrics
    assert song["slide_breaks"] == {"chorus": [2]}


def test_manual_break_forces_new_slide_in_deck_and_live():
    song = app.normalize_worship_song({
        "title": "T",
        "parts": {"chorus": ["Line one", "Line two", "Line three", "Line four"]},
        "slide_breaks": {"chorus": [1]},
        "arrangement": ["chorus"],
    })
    chunks = [chunk["lines"] for _, chunk in app._iter_worship_lyric_chunks(song)]
    assert chunks == [["Line one"], ["Line two", "Line three", "Line four"]]


def test_song_without_breaks_is_unchanged():
    song = app.normalize_worship_song({
        "title": "T",
        "parts": {"chorus": ["Line one", "Line two", "Line three", "Line four"]},
        "arrangement": ["chorus"],
    })
    assert song["slide_breaks"] == {}
    chunks = [chunk["lines"] for _, chunk in app._iter_worship_lyric_chunks(song)]
    assert chunks == [["Line one", "Line two", "Line three", "Line four"]]


def test_video_times_accept_seconds_or_minutes():
    assert app._parse_video_seconds("90") == 90
    assert app._parse_video_seconds("1:30") == 90
    assert app._parse_video_seconds("1:02:03") == 3723
    assert app._parse_video_seconds("") == 0
    assert app._parse_video_seconds("abc") == 0
    assert app._parse_video_seconds("-5") == 0



PSALM_23 = (
    "The Lord is my shepherd; I shall not want. He makes me lie down in green pastures. "
    "He leads me beside still waters. He restores my soul. He leads me in paths of righteousness "
    "for his name's sake. Even though I walk through the valley of the shadow of death, I will fear "
    "no evil, for you are with me; your rod and your staff, they comfort me. You prepare a table "
    "before me in the presence of my enemies; you anoint my head with oil; my cup overflows. Surely "
    "goodness and mercy shall follow me all the days of my life, and I shall dwell in the house of "
    "the Lord forever."
)


def _scripture_slides(text):
    lines, breaks = app._worship_scripture_slide_lines(text)
    item = app.normalize_worship_song({
        "title": "Psalm 23", "type": "scripture", "parts": {"reading": lines},
        "slide_breaks": {"reading": breaks}, "arrangement": ["reading"],
    })
    return lines, [chunk for _, chunk in app._iter_worship_lyric_chunks(item)]


def test_scripture_slides_end_on_sentence_boundaries_and_keep_exact_text():
    lines, slides = _scripture_slides(PSALM_23)
    assert " ".join(lines) == PSALM_23
    for slide in slides:
        assert slide["lines"][-1].rstrip().endswith((".", "!", "?", ";", ":", ",")), slide["lines"]


def test_scripture_lines_fit_without_rewrapping_in_powerpoint():
    _, slides = _scripture_slides(PSALM_23)
    for slide in slides:
        assert slide["font_size"] >= 36
        for line in slide["lines"]:
            assert app._lyric_width_units(line) * slide["font_size"] <= 874


def test_one_huge_sentence_still_splits_safely():
    text = ("word " * 120).strip() + "."
    lines, breaks = app._worship_scripture_slide_lines(text)
    assert " ".join(lines) == text
    assert breaks  # split into several slides


def test_local_today_uses_browser_time_zone_cookie():
    from datetime import datetime
    from unittest import mock

    fixed = datetime(2026, 10, 31, 1, 30, tzinfo=app.timezone.utc)  # Fri 9:30 PM in New York
    with mock.patch.object(app, "datetime", wraps=datetime) as fake:
        fake.now.side_effect = lambda tz=None: fixed.astimezone(tz) if tz else fixed
        with app.app.test_request_context("/", headers={"Cookie": "fs_tz=America/New_York"}):
            assert app._worship_local_today() == "2026-10-30"
        with app.app.test_request_context("/"):
            assert app._worship_local_today() == "2026-10-31"
        with app.app.test_request_context("/", headers={"Cookie": "fs_tz=Not/AZone"}):
            assert app._worship_local_today() == "2026-10-31"


def test_quick_scripture_accepts_whole_chapter_and_refuses_to_truncate():
    from unittest import mock

    with mock.patch("faithsparks.services.scripture.fetch_verse_text", return_value="The Lord is my shepherd."):
        item = app._build_quick_worship_scripture("Psalm 23", "web")
    assert item["title"] == "Psalm 23"
    with mock.patch("faithsparks.services.scripture.fetch_verse_text", return_value="word " * 2000):
        try:
            app._build_quick_worship_scripture("Psalm 119", "kjv")
        except ValueError as exc:
            assert "too long" in str(exc)
        else:
            raise AssertionError("long passage should be refused, not truncated")


def test_public_domain_miss_says_check_reference():
    from unittest import mock

    with mock.patch("faithsparks.services.scripture.fetch_verse_text", return_value=None):
        try:
            app._build_quick_worship_scripture("John 99:1", "kjv")
        except RuntimeError as exc:
            assert "Check the reference" in str(exc) and "Try WEB or KJV" not in str(exc)


def test_review_page_shows_repeat_count():
    import re
    from flask import render_template

    song = _song(["verse1", "chorus", "chorus"])
    with app.app.test_request_context("/worship/deck/review", method="POST"):
        slides = app._build_worship_mobile_slides([song])
        html = render_template(
            "worship_deck_review.html", selected_items=[song], slides=slides, findings=[],
            crowded_count=0, song_order="test-song", review_fingerprint="x",
        )
    html = re.sub(r"<script\b.*?</script>", "", html, flags=re.S)
    assert "×2 (repeat this slide)" in html
    assert html.count("(repeat this slide)") == 1
