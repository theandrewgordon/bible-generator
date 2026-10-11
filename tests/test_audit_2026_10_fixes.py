"""Regression tests for the 2026-10-11 audit fixes (slide builder + live + app config)."""

import io
import re
import tempfile
import time
import unittest
import zipfile
from unittest import mock

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import app
from faithsparks.services import rate_limit


def _deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def _slide_xml(prs) -> str:
    buffer = io.BytesIO()
    prs.save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        return "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )


class SlideRenderingTests(unittest.TestCase):
    def test_multiline_text_uses_real_paragraphs(self):
        prs = _deck()
        app.create_content_slide(prs, ["Line one", "Line two", "Line three"], "song", None)
        xml = _slide_xml(prs)
        texts = re.findall(r"<a:t>(.*?)</a:t>", xml, flags=re.S)
        self.assertFalse(any("\n" in text for text in texts), "raw newline inside <a:t>")
        self.assertIn("Line two", texts)
        self.assertGreaterEqual(xml.count('algn="ctr"'), 3)

    def test_cover_picture_crops_instead_of_stretching(self):
        prs = _deck()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with tempfile.TemporaryDirectory() as tmp:
            wide = f"{tmp}/wide.png"
            tall = f"{tmp}/tall.png"
            Image.new("RGB", (1536, 1024), "navy").save(wide)  # 3:2
            Image.new("RGB", (600, 1200), "navy").save(tall)  # portrait photo
            pic = app._add_cover_picture(slide, wide, 0, 0, 13.333, 7.5)
            self.assertAlmostEqual(pic.crop_left, 0.0)
            self.assertGreater(pic.crop_top, 0.05)  # 3:2 is taller than 16:9
            pic = app._add_cover_picture(slide, tall, 0, 0, 13.333, 7.5)
            self.assertGreater(pic.crop_top, 0.3)
            pic = app._add_cover_picture(slide, tall, 0, 0, 7.0, 7.5)
            self.assertAlmostEqual(pic.crop_left, 0.0)

    def test_backgrounds_resolve_to_small_16x9_jpegs(self):
        for name in app._load_bg_config():
            path, _ = app._resolve_bg("song", name)
            self.assertIsNotNone(path, name)
            self.assertEqual(path.suffix, ".jpg")
            self.assertLess(path.stat().st_size, 600_000)
            with Image.open(path) as image:
                self.assertAlmostEqual(image.width / image.height, 16 / 9, places=2)

    def test_long_divider_titles_shrink(self):
        self.assertEqual(app._divider_title_font_size("Holy Forever"), 60)
        self.assertLess(app._divider_title_font_size("Test Hymn With A Rather Long Title That Keeps Going"), 48)

    def test_footer_is_plain_text_not_hyperlink(self):
        prs = _deck()
        app.create_divider_slide(prs, "Title", "Artist", "G")
        self.assertNotIn("hlinkClick", _slide_xml(prs))

    def test_continuation_check_uses_whole_words(self):
        for line in ("Jesus my Savior", "He is coming again", "Washed away my sin", "Our Creator"):
            self.assertFalse(app._line_continues_to_next(line), line)
        for line in ("Lift your hands and", "Come to", "Holy, holy,"):
            self.assertTrue(app._line_continues_to_next(line), line)


class _FakeSnap:
    def __init__(self, data):
        self._data = data

    @property
    def exists(self):
        return self._data is not None

    def to_dict(self):
        return dict(self._data) if self._data is not None else None


class _FakeDoc:
    def __init__(self, store, key):
        self.store, self.key, self.calls = store, key, []

    def get(self, field_paths=None, transaction=None):
        self.calls.append(tuple(field_paths) if field_paths else None)
        data = self.store.get(self.key)
        if data is not None and field_paths:
            data = {k: v for k, v in data.items() if k in field_paths}
        return _FakeSnap(data)


class _FakeTxn:
    def __init__(self, store):
        self.store = store

    def set(self, ref, fields, merge=False):
        self.store[ref.key].update(fields)


class LiveWorshipFirestorePathTests(unittest.TestCase):
    def setUp(self):
        self.store = {
            "sess": {
                "id": "sess", "scope": "grace", "slide_count": 3, "current_index": 0,
                "revision": 0, "deck_revision": 0, "expires_epoch": time.time() + 3600,
                "slides": [{"kind": "lyric", "lines": ["a", "b"], "font_size": 48}] * 3,
            }
        }
        self.doc = _FakeDoc(self.store, "sess")
        fake_db = mock.Mock()
        fake_db.transaction.side_effect = lambda: _FakeTxn(self.store)
        self.patches = [
            mock.patch.object(app, "db", fake_db),
            mock.patch.object(app, "_worship_live_ref_for_scope",
                              lambda scope: mock.Mock(document=lambda sid: self.doc)),
            mock.patch.object(app.firestore, "transactional", lambda fn: fn),
        ]
        for patch in self.patches:
            patch.start()

    def tearDown(self):
        for patch in reversed(self.patches):
            patch.stop()

    def test_poll_reads_only_state_fields(self):
        live = app._get_worship_live_session("grace", "a" * 24, include_slides=False)
        self.assertNotIn("slides", live)
        self.assertEqual(self.doc.calls[-1], app._WORSHIP_LIVE_STATE_FIELDS)

    def test_navigation_does_not_read_or_rewrite_slides(self):
        state = app._update_worship_live_session("grace", "sess", "next")
        self.assertEqual(state["current_index"], 1)
        self.assertEqual(self.doc.calls[-1], app._WORSHIP_LIVE_STATE_FIELDS)
        self.assertEqual(len(self.store["sess"]["slides"]), 3)

    def test_navigation_clamps_without_slides_loaded(self):
        for _ in range(5):
            state = app._update_worship_live_session("grace", "sess", "next")
        self.assertEqual(state["current_index"], 2)

    def test_deck_edit_reads_and_writes_slides(self):
        app._update_worship_live_session("grace", "sess", "split_slide", 0, split_after=1)
        self.assertIsNone(self.doc.calls[-1])
        self.assertEqual(len(self.store["sess"]["slides"]), 4)
        self.assertEqual(self.store["sess"]["slide_count"], 4)


class AppConfigTests(unittest.TestCase):
    def test_request_body_cap_is_set(self):
        self.assertEqual(app.app.config["MAX_CONTENT_LENGTH"], 30 * 1024 * 1024)

    def test_oversized_request_gets_friendly_413(self):
        client = app.app.test_client()
        response = client.post("/worship/presentation/import", data=b"x" * (31 * 1024 * 1024),
                               content_type="application/octet-stream")
        self.assertEqual(response.status_code, 413)

    def test_local_only_rate_limit_skips_firestore(self):
        rate_limit.reset_memory_limits()
        fake_db = mock.Mock()
        with mock.patch.object(rate_limit, "db", fake_db):
            result = rate_limit.check_rate_limit("t", "k", limit=1, window_seconds=60, local_only=True)
            self.assertTrue(result.allowed)
            self.assertFalse(rate_limit.check_rate_limit("t", "k", limit=1, window_seconds=60, local_only=True).allowed)
        fake_db.collection.assert_not_called()
        rate_limit.reset_memory_limits()


if __name__ == "__main__":
    unittest.main()
